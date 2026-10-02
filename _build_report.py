from pathlib import Path
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Academic-RAG")
OUT = ROOT / "docs" / "Academic_RAG_Project_Report.docx"
WORK = Path(tempfile.gettempdir()) / "academic_rag_report_build"
WORK.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)

BLACK = "000000"
INK = "283445"
MUTED = "5F6D7D"
NAVY = "203A5C"
TEAL = "187D83"
PALE_BLUE = "EAF1F7"
PALE_TEAL = "E8F4F2"
PALE_GRAY = "F4F6F8"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
FONT_NAME = "Arial"


def load_font(size, bold=False):
    paths = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
    ]
    for path in paths:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def centered_multiline(draw, box, text, font, fill, spacing=8):
    x1, y1, x2, y2 = box
    draw.multiline_text(
        ((x1 + x2) / 2, (y1 + y2) / 2),
        text,
        font=font,
        fill=fill,
        anchor="mm",
        align="center",
        spacing=spacing,
    )


def draw_arrow(draw, start, end, color=(91, 109, 127), width=6, head=17):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - direction * head, y2 - head), (x2 - direction * head, y2 + head)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - head, y2 - direction * head), (x2 + head, y2 - direction * head)]
    draw.polygon(points, fill=color)


def make_architecture_diagram(path):
    width, height = 1800, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    lane_font = load_font(29, True)
    node_font = load_font(27, True)
    small_font = load_font(21, False)

    draw.text((82, 58), "DOCUMENT INGESTION", font=lane_font, fill=f"#{NAVY}")
    draw.text((82, 468), "QUESTION ANSWERING", font=lane_font, fill=f"#{NAVY}")
    draw.line((82, 101, 1715, 101), fill=f"#{BORDER}", width=2)
    draw.line((82, 511, 1715, 511), fill=f"#{BORDER}", width=2)

    xs = [90, 425, 760, 1095, 1430]
    box_w, box_h = 250, 170
    top_y, bottom_y = 140, 560
    top_nodes = [
        "PDF upload\nStreamlit interface",
        "Validate and\nextract page text\npypdf",
        "Split into chunks\n800 characters\n150 overlap",
        "Create vectors\nGemini embeddings",
        "Store text and vectors\nChroma with source\nand page metadata",
    ]
    bottom_nodes = [
        "Student asks\na question\nStreamlit chat",
        "Create query vector\nGemini embeddings",
        "Search closest chunks\nChroma returns\nup to three",
        "Generate from evidence\nGemini language model",
        "Show answer with\nsource and page\nStreamlit interface",
    ]
    top_colors = [PALE_BLUE, PALE_BLUE, PALE_BLUE, PALE_TEAL, PALE_TEAL]
    bottom_colors = [PALE_BLUE, PALE_TEAL, PALE_TEAL, PALE_BLUE, PALE_BLUE]

    for y, nodes, colors in ((top_y, top_nodes, top_colors), (bottom_y, bottom_nodes, bottom_colors)):
        for x, label, fill in zip(xs, nodes, colors):
            draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=20, fill=f"#{fill}", outline=f"#{NAVY}", width=3)
            centered_multiline(draw, (x + 14, y + 12, x + box_w - 14, y + box_h - 12), label, node_font, f"#{INK}", 7)

    for index in range(4):
        draw_arrow(draw, (xs[index] + box_w + 12, top_y + box_h // 2), (xs[index + 1] - 14, top_y + box_h // 2))
        draw_arrow(draw, (xs[index] + box_w + 12, bottom_y + box_h // 2), (xs[index + 1] - 14, bottom_y + box_h // 2))

    # The indexed vectors and metadata created above feed the later similarity search.
    store_center_x = xs[4] + box_w // 2
    search_center_x = xs[2] + box_w // 2
    connector_y = 445
    draw.line([(store_center_x, top_y + box_h + 7), (store_center_x, connector_y)], fill=f"#{TEAL}", width=5)
    draw.line([(store_center_x, connector_y), (search_center_x, connector_y)], fill=f"#{TEAL}", width=5)
    draw_arrow(draw, (search_center_x, connector_y), (search_center_x, bottom_y - 13), color=tuple(bytes.fromhex(TEAL)), width=5, head=16)
    draw.text((1125, connector_y - 32), "stored vectors and metadata", font=small_font, fill=f"#{TEAL}")

    image.save(path, dpi=(220, 220))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=105, start=115, bottom=105, end=115):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def format_cell_paragraph(paragraph, text, *, bold=False, color=INK, size=9.5, align=None):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers, rows, widths, body_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=120, bottom=120)
        cell.text = ""
        format_cell_paragraph(cell.paragraphs[0], label, bold=True, color=WHITE, size=9.4)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else PALE_GRAY)
            cell.text = ""
            format_cell_paragraph(
                cell.paragraphs[0],
                value,
                bold=(index == 0),
                color=INK,
                size=body_size,
            )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.45
    return table


def set_run_font(run, name=FONT_NAME, size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_body(doc, text, *, after=6, before=0, size=11, keep=False):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.widow_control = True
    paragraph.paragraph_format.keep_together = keep
    set_run_font(paragraph.add_run(text), size=size)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(text)
    return paragraph


def add_step(doc, number, lead, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.first_line_indent = Inches(-0.08)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.06
    run = paragraph.add_run(f"{number}  {lead} ")
    set_run_font(run, size=10.6, bold=True, color=NAVY)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.6)
    return paragraph


def add_bullet(doc, lead, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.13)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.04
    run = paragraph.add_run(lead)
    set_run_font(run, size=10.5, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_page_number(paragraph):
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.34)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = FONT_NAME
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.line_spacing = 1.0

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = FONT_NAME
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(5)

    for style_name, size, before, after in (("Heading 1", 18, 0, 9), ("Heading 2", 12.5, 7, 4)):
        style = doc.styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.95), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(header.add_run("ACADEMIC RAG STUDY ASSISTANT"), size=8.2, color=BLACK, bold=True)
    set_run_font(header.add_run("\tPROJECT REPORT"), size=8.2, color=BLACK)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    set_run_font(footer.add_run("B Tech CSE Mini Project     "), size=8.5, color=MUTED)
    add_page_number(footer)


def page_break(doc):
    doc.add_page_break()


def main():
    diagram_path = WORK / "academic_rag_architecture.png"
    make_architecture_diagram(diagram_path)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Academic RAG Study Assistant Project Report"
    doc.core_properties.subject = "Design, architecture, implementation, testing, and future scope"
    doc.core_properties.keywords = "Academic RAG, Gemini, Streamlit, Chroma, project report"

    # Page one introduces the project and explains the problem it addresses.
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.keep_with_next = True
    title.add_run("Academic RAG Study Assistant\nProject Report")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Document grounded question answering for academic study")
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(16)
    set_run_font(meta.add_run("B Tech Computer Science and Engineering Mini Project"), size=10, color=TEAL, bold=True)

    add_heading(doc, "2 Abstract", 1)
    add_body(
        doc,
        "Students often search lecture notes and subject PDFs manually. This takes time, and answers from general search may not be supported by the material assigned for a course. Academic RAG addresses this problem by letting a student upload a searchable PDF and ask questions in natural language. The system retrieves passages from the uploaded material, asks a language model to answer from that evidence, and displays the source file and page references so the student can verify the result.",
        after=7,
    )
    add_body(
        doc,
        "The application extracts page text with pypdf, divides it into overlapping chunks, requests embeddings from Google's Gemini API, and stores each chunk with its source and page metadata in Chroma. For a question, the application embeds the query, retrieves up to three relevant passages, and sends those passages to Gemini for grounded answer generation. Streamlit provides PDF upload, chat, library statistics, session history, and a light or dark appearance. The current code also validates PDFs and reports common configuration, API, and storage errors.",
        after=7,
    )
    add_body(
        doc,
        "A development run with the supplied three page test PDF reported 10 chunks and 10 stored records. The end to end script produced an answer with source page references. This report describes the current local prototype; deployment, per user data isolation, durable cloud storage, and measured load performance remain future work.",
        after=10,
    )
    add_heading(doc, "Project Overview", 2)
    add_table(
        doc,
        ["Project element", "Current implementation"],
        [
            ("Input", "Searchable academic PDFs uploaded through the Streamlit interface"),
            ("Core process", "Extract, split, embed, store, retrieve, and generate an evidence based answer"),
            ("Output", "A concise answer with the source PDF name and page references"),
            ("Status", "Working local prototype with manual development checks"),
        ],
        [1.45, 5.53],
        body_size=9.6,
    )

    # Page two explains the technologies and their roles.
    page_break(doc)
    add_heading(doc, "3 Technology Stack", 1)
    add_body(
        doc,
        "The project uses Python to connect the interface, PDF processing, vector search, and Gemini API. Each main responsibility lives in its own module, which makes the data path easier to follow and individual parts easier to maintain.",
        after=8,
    )
    add_table(
        doc,
        ["Layer", "Technology", "Role in this project"],
        [
            ("User interface", "Streamlit", "PDF uploader, chat, indexed library view, session history, and light or dark mode."),
            ("PDF extraction", "pypdf", "Checks and reads PDF files while retaining page numbers; scanned image PDFs need OCR, which is not implemented."),
            ("Chunking", "Python module", "Splits page text into 800 character chunks with 150 characters of overlap."),
            ("Embeddings", "Google Gemini API\ngemini-embedding-2", "Converts each chunk and question into a numeric vector for semantic matching."),
            ("Vector store", "ChromaDB", "Persists vectors, text, and source/page metadata locally in the academic_notes collection; returns up to three matches."),
            ("Answer generation", "Google Gemini API\ngemini-3.5-flash-lite", "Uses retrieved passages as context and is prompted to answer from that evidence and cite sources."),
            ("Configuration and errors", "python-dotenv\nStreamlit secrets", "Loads the API key locally or from Streamlit Cloud settings and maps expected failures to readable messages."),
        ],
        [1.18, 1.72, 4.08],
        body_size=9.15,
    )
    add_heading(doc, "Technology Choices", 2)
    add_body(
        doc,
        "Streamlit keeps the interface in Python. Gemini provides managed embedding and answer generation, so the project does not need to download a large language model. Chroma runs in persistent local mode for the current build. The API key stays outside the source code: local work uses .env, while the Gemini client can also read st.secrets when configured for Streamlit Cloud.",
        after=6,
    )
    add_body(
        doc,
        "The local database and uploaded files are not durable cloud storage. A deployed multi user version will need a persistent hosted store and a way to separate each user's documents.",
        after=0,
    )

    # Page three combines the architecture diagram with a plain language walkthrough.
    page_break(doc)
    add_heading(doc, "4 System Flow and Architecture", 1)
    add_body(
        doc,
        "The system has an ingestion path for PDFs and a question answering path for student queries. Both paths meet at the vector store: ingestion writes searchable passages, while retrieval reads the best matching passages for a question.",
        after=5,
    )
    figure = doc.add_picture(str(diagram_path), width=Inches(6.95))
    figure._inline.docPr.set("descr", "Diagram showing PDF extraction and embedding into Chroma, followed by question embedding, top three retrieval, Gemini answer generation, and source page display.")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)
    set_run_font(caption.add_run("Figure 1  Academic RAG document ingestion and question answering paths"), size=8.7, color=MUTED, italic=True)

    add_step(doc, "1", "Upload and extract", "The user uploads a PDF. The application checks the file and extracts text one page at a time while preserving the page number.")
    add_step(doc, "2", "Split and index", "The splitter creates overlapping chunks. Gemini embeds each chunk, and Chroma stores the vector, text, filename, and page.")
    add_step(doc, "3", "Retrieve evidence", "The application embeds the question and asks Chroma for up to three semantically similar passages.")
    add_step(doc, "4", "Generate and cite", "Gemini receives the question and retrieved passages. The answer is shown with source and page details; chat messages remain in the active Streamlit session.")

    # Page four documents the actual on-disk organization.
    page_break(doc)
    add_heading(doc, "5 Directory Structure", 1)
    add_body(
        doc,
        "The source is organized around the application entry point and focused modules under src. The tree below reflects the current project and its local data folders.",
        after=7,
    )
    tree = [
        "Academic-RAG/",
        "├── app.py                                  Streamlit user interface",
        "├── requirements.txt                        Python dependencies",
        "├── .env.example                            API key placeholder",
        "├── .gitignore                               excludes local secrets and data",
        "├── .streamlit/",
        "│   └── config.toml                          default theme configuration",
        "├── src/",
        "│   ├── pdf_loader.py                        PDF checks and page extraction",
        "│   ├── text_splitter.py                     overlapping text chunks",
        "│   ├── embeddings.py                        Gemini vector requests",
        "│   ├── vector_store.py                      Chroma storage and search",
        "│   ├── rag_chain.py                         evidence grounded answer generation",
        "│   ├── gemini_client.py                     API client and secret lookup",
        "│   └── errors.py                            expected error messages",
        "├── test_search.py                          manual retrieval script",
        "├── test_rag.py                             manual end to end script",
        "└── data/",
        "    ├── uploads/                            local uploaded PDFs",
        "    ├── chroma_db/                          local vector database",
        "    └── Academic_RAG_Test_Document_3_Pages.pdf",
    ]
    for line in tree:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.07)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_together = True
        set_run_font(paragraph.add_run(line), name="Consolas", size=8.25, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    add_heading(doc, "Module Responsibilities", 2)
    add_body(
        doc,
        "app.py coordinates the user interaction. pdf_loader.py and text_splitter.py prepare content; embeddings.py and gemini_client.py call the remote model service; vector_store.py manages Chroma retrieval; and rag_chain.py builds the grounded response. errors.py defines the safe exceptions surfaced by the interface. The two test_*.py files are manual smoke scripts rather than a formal automated test suite.",
        after=7,
    )
    add_body(
        doc,
        "The actual .env file contains a private API key and is excluded by .gitignore. The uploads and Chroma folders are also local working data. They are omitted from Git so the repository does not expose credentials or accidentally publish student documents.",
        after=0,
    )

    # Page five reports observed checks and separates evidence from unmeasured claims.
    page_break(doc)
    add_heading(doc, "6 Testing and Performance", 1)
    add_body(
        doc,
        "The current evidence comes from manual command line checks on the supplied three page PDF. These checks demonstrate the main pipeline, but they are not a substitute for automated regression tests or a load benchmark.",
        after=7,
    )
    add_table(
        doc,
        ["Check", "Observed development result"],
        [
            ("PDF ingestion and chunking", "The sample PDF produced 10 text chunks; the save command reported 10 stored Chroma records."),
            ("Embedding output", "A console check returned one vector containing 3,072 numeric values for a sample text."),
            ("Semantic retrieval", "The retrieval script searches for a question and prints matching passages with their source and page."),
            ("End to end RAG", "The answer script returned a response to a RAG question and displayed citations to pages 2 and 3 of the sample PDF."),
        ],
        [1.65, 5.33],
        body_size=9.5,
    )
    add_heading(doc, "Performance Status", 2)
    add_body(
        doc,
        "No ingestion time, query latency, throughput, or concurrent user benchmark has been recorded, so this report makes no speed or capacity claim. The current implementation requests an embedding for each chunk sequentially. A question uses one query embedding, a local Chroma search of up to three results, and one Gemini generation request. Actual response time therefore depends on document size, network conditions, and Gemini service availability.",
        after=8,
    )
    add_heading(doc, "7 Conclusion and Future Scope", 1)
    add_body(
        doc,
        "Academic RAG demonstrates a complete retrieval augmented generation workflow for academic PDFs: it extracts and chunks text, indexes embeddings, retrieves relevant evidence, and presents a sourced answer. The manual checks confirm that the local prototype can process a small test PDF and answer a document based question. Its present scope is a single shared local library, and deployment has intentionally been postponed.",
        after=5,
    )
    add_bullet(doc, "User accounts and isolation. ", "Give each student a separate document collection and chat history before opening the app to multiple users.")
    add_bullet(doc, "Durable storage. ", "Move PDFs and vectors to hosted object and vector storage so data survives application restarts.")
    add_bullet(doc, "Broader document support. ", "Add OCR for scanned PDFs and test extraction against tables, diagrams, and difficult layouts.")
    add_bullet(doc, "Measured quality and speed. ", "Add repeatable automated tests, retrieval and citation quality checks, API retry controls, and latency or concurrency benchmarks.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
